"""
app.py
Audio / Video → Executive Summary + Minutes of Meeting + Feedback Email
Transasia Biomedicals Ltd.
"""

import os
import re
import time
import tempfile
import subprocess
from io import BytesIO
from pathlib import Path

import streamlit as st
import anthropic
from docx import Document
from groq import Groq
from dotenv import load_dotenv

load_dotenv()

# ── API keys ──────────────────────────────────────────────────
# On Streamlit Cloud: set these in App Settings → Secrets.
# Locally: set them in .env.
ANTHROPIC_API_KEY = (
    st.secrets.get("ANTHROPIC_API_KEY")
    or os.environ.get("ANTHROPIC_API_KEY", "")
)
GROQ_API_KEY = (
    st.secrets.get("GROQ_API_KEY")
    or os.environ.get("GROQ_API_KEY", "")
)

# ── Metadata defaults ─────────────────────────────────────────
# On Streamlit Cloud: optionally set these in Secrets to pre-fill the sidebar.
# Locally: set them in .env. Falls back to empty/generic placeholders.
def _secret(key: str, fallback: str = "") -> str:
    return st.secrets.get(key) or os.environ.get(key, fallback)

_DEF_SENDER_NAME  = _secret("SENDER_NAME",       "Your Name")
_DEF_SENDER_TITLE = _secret("SENDER_TITLE",      "Your Title")
_DEF_COMPANY      = _secret("COMPANY",           "Your Company")
_DEF_RECIPIENT    = _secret("RECIPIENT_NAME",    "Recipient Name")
_DEF_SUBJECT      = _secret("EMAIL_SUBJECT",     "Meeting / Session Feedback")
_DEF_CONTEXT      = _secret("FEEDBACK_CONTEXT",  "Brief context of the recording")

# ── Page config ───────────────────────────────────────────────
st.set_page_config(
    page_title="Audio/Video to Text Data Conversion",
    page_icon="🎙️",
    layout="centered",
)

st.title("Audio/Video to Text Data Conversion")
st.caption(
    "Converts audio/video recordings to verbatim Summary, Minutes of Meeting, "
    "and Feedback Email. Only what was said is captured — nothing is added."
)

# ── Sidebar ───────────────────────────────────────────────────
with st.sidebar:
    st.header("Configuration")
    language = st.selectbox("Transcription language", ["en", "hi", "auto"], index=0)
    chunk_minutes = st.slider(
        "Chunk size (minutes)", min_value=10, max_value=30, value=20, step=5,
        help="Audio is split into chunks of this size. Smaller = safer for rate limits.",
    )
    st.divider()
    st.subheader("Generate")
    do_summary = st.checkbox("📋 Summary",            value=True)
    do_mom     = st.checkbox("📝 Minutes of Meeting", value=True)
    do_email   = st.checkbox("📧 Feedback Email",     value=True)
    if not any([do_summary, do_mom, do_email]):
        st.warning("Select at least one output.")
    st.divider()
    st.subheader("Document metadata")
    sender_name      = st.text_input("Your name",     _DEF_SENDER_NAME)
    sender_title     = st.text_input("Your title",    _DEF_SENDER_TITLE)
    company          = st.text_input("Company",       _DEF_COMPANY)
    recipient_name   = st.text_input("Recipient",     _DEF_RECIPIENT)
    email_subject    = st.text_input("Email subject", _DEF_SUBJECT)
    feedback_context = st.text_input("Context",       _DEF_CONTEXT)

# ── Helpers ───────────────────────────────────────────────────
def fmt_time(seconds: float) -> str:
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def _add_inline(para, text: str):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if part.startswith("**") and part.endswith("**"):
            para.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            para.add_run(part[1:-1]).italic = True
        elif part:
            para.add_run(part)


def markdown_to_docx(markdown_text: str, title: str = "") -> bytes:
    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    lines = markdown_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.lstrip()
        indent   = len(line) - len(stripped)

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("|"):
            # Table block
            tbl_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            rows = [
                [c.strip() for c in row.strip().strip("|").split("|")]
                for row in tbl_lines
            ]
            rows = [r for r in rows if not all(re.fullmatch(r"[-: ]+", c) for c in r)]
            if rows:
                ncols = max(len(r) for r in rows)
                tbl = doc.add_table(rows=len(rows), cols=ncols)
                tbl.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, cell_text in enumerate(row):
                        cell = tbl.cell(ri, ci)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        if ri == 0:
                            p.add_run(cell_text).bold = True
                        else:
                            _add_inline(p, cell_text)
            continue
        elif re.match(r"^(\s*)[-*] ", line):
            # Bullet — depth based on indent level
            if indent >= 4:
                style = "List Bullet 3"
            elif indent >= 2:
                style = "List Bullet 2"
            else:
                style = "List Bullet"
            content = re.sub(r"^(\s*)[-*] ", "", line).strip()
            p = doc.add_paragraph(style=style)
            _add_inline(p, content)
        elif re.match(r"^(\s*)\d+\. ", line):
            if indent >= 4:
                style = "List Number 3"
            elif indent >= 2:
                style = "List Number 2"
            else:
                style = "List Number"
            content = re.sub(r"^(\s*)\d+\. ", "", line).strip()
            p = doc.add_paragraph(style=style)
            _add_inline(p, content)
        elif line.strip():
            p = doc.add_paragraph()
            _add_inline(p, line.strip())
        i += 1

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# ── Session state ─────────────────────────────────────────────
_defaults = {
    "stage":      "upload",
    "chunks":     None,
    "stem":       None,
    "file_id":    None,
    "transcript": None,
    "summary":    None,
    "mom":        None,
    "email":      None,
}
for k, v in _defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v

# ── File upload ───────────────────────────────────────────────
uploaded = st.file_uploader(
    "Upload a video or audio file",
    type=["mkv","mp4","avi","mov","webm","mp3","m4a","wav","ogg","flac","aac"],
)

if not uploaded:
    st.session_state.stage  = "upload"
    st.session_state.chunks = None
    st.info("Upload a file above to get started.")
    st.stop()

# Reset if a different file is uploaded
file_id = f"{uploaded.name}_{uploaded.size}"
if st.session_state.file_id != file_id:
    for k, v in _defaults.items():
        st.session_state[k] = v
    st.session_state.file_id = file_id

if not any([do_summary, do_mom, do_email]):
    st.warning("Please select at least one output in the sidebar.")
    st.stop()

# ═══════════════════════════════════════════════════════════════
# STAGE 1 — Extract & split audio
# ═══════════════════════════════════════════════════════════════
if st.session_state.stage == "upload":
    if st.button("▶ Extract & Split Audio", type="primary"):
        with tempfile.TemporaryDirectory() as tmp:
            tmp        = Path(tmp)
            input_path = tmp / uploaded.name
            audio_path = tmp / "audio.mp3"
            input_path.write_bytes(uploaded.read())

            with st.status("Step 1 — Extracting audio …") as s:
                res = subprocess.run(
                    ["ffmpeg", "-y", "-i", str(input_path),
                     "-vn", "-acodec", "libmp3lame",
                     "-ar", "16000", "-ac", "1", "-b:a", "32k",
                     str(audio_path)],
                    capture_output=True, text=True,
                )
                if res.returncode != 0:
                    st.error(f"FFmpeg failed:\n{res.stderr}")
                    st.stop()
                s.update(label="Step 1 — Audio extracted", state="complete")

            probe = subprocess.run(
                ["ffprobe", "-v", "error",
                 "-show_entries", "format=duration",
                 "-of", "default=noprint_wrappers=1:nokey=1",
                 str(audio_path)],
                capture_output=True, text=True,
            )
            total_dur = float(probe.stdout.strip())

            chunk_sec = chunk_minutes * 60
            with st.status(f"Step 2 — Splitting into {chunk_minutes}-min chunks …") as s:
                subprocess.run(
                    ["ffmpeg", "-y", "-i", str(audio_path),
                     "-f", "segment", "-segment_time", str(chunk_sec),
                     "-c", "copy", str(tmp / "chunk_%03d.mp3")],
                    capture_output=True,
                )
                chunk_files = sorted(tmp.glob("chunk_*.mp3"))
                chunks = []
                for idx, cf in enumerate(chunk_files):
                    start   = idx * chunk_sec
                    end     = min((idx + 1) * chunk_sec, total_dur)
                    size_mb = cf.stat().st_size / 1_000_000
                    chunks.append({
                        "idx":      idx,
                        "start":    start,
                        "end":      end,
                        "duration": end - start,
                        "label":    (
                            f"Chunk {idx+1}:  {fmt_time(start)} – {fmt_time(end)}"
                            f"  ({(end-start)/60:.1f} min,  {size_mb:.1f} MB)"
                        ),
                        "bytes":    cf.read_bytes(),
                        "name":     cf.name,
                    })
                st.session_state.chunks = chunks
                st.session_state.stem   = Path(uploaded.name).stem
                s.update(label=f"Step 2 — Split into {len(chunks)} chunk(s)", state="complete")

        st.session_state.stage = "select"
        st.rerun()

# ═══════════════════════════════════════════════════════════════
# STAGE 2 — Select chunks to transcribe
# ═══════════════════════════════════════════════════════════════
elif st.session_state.stage == "select":
    chunks = st.session_state.chunks

    st.subheader(f"Select chunks to transcribe  ({len(chunks)} total)")

    for c in chunks:
        if f"chunk_{c['idx']}" not in st.session_state:
            st.session_state[f"chunk_{c['idx']}"] = True

    col1, col2 = st.columns(2)
    if col1.button("☑ Select All"):
        for c in chunks:
            st.session_state[f"chunk_{c['idx']}"] = True
        st.rerun()
    if col2.button("☐ Deselect All"):
        for c in chunks:
            st.session_state[f"chunk_{c['idx']}"] = False
        st.rerun()

    st.divider()
    for chunk in chunks:
        st.checkbox(chunk["label"], key=f"chunk_{chunk['idx']}")

    selected    = [c for c in chunks if st.session_state.get(f"chunk_{c['idx']}", True)]
    total_sel_s = sum(c["duration"] for c in selected)

    st.divider()
    if not selected:
        st.warning("Select at least one chunk.")
        st.stop()

    GROQ_HOURLY_LIMIT = 7200
    if total_sel_s > GROQ_HOURLY_LIMIT:
        hrs = total_sel_s / GROQ_HOURLY_LIMIT
        st.warning(
            f"**Selected: {total_sel_s/60:.0f} min of audio.**  "
            f"Groq free tier allows 120 min/hour — this will span "
            f"~{hrs:.1f} hour-windows. The app will pause with a countdown "
            f"if a rate-limit is hit."
        )
    else:
        st.info(f"Selected: {total_sel_s/60:.0f} min — fits within Groq's hourly limit.")

    if st.button("▶ Transcribe & Generate", type="primary"):

        # ── Transcribe each selected chunk ────────────────────
        groq_client = Groq(api_key=GROQ_API_KEY)
        parts = []

        for i, chunk in enumerate(selected):
            short = f"Chunk {chunk['idx']+1} ({fmt_time(chunk['start'])}–{fmt_time(chunk['end'])})"
            with st.status(f"Transcribing {short}  [{i+1}/{len(selected)}] …") as s:
                tf = tempfile.NamedTemporaryFile(suffix=".mp3", delete=False)
                tf.write(chunk["bytes"])
                tf.close()
                try:
                    for attempt in range(3):
                        try:
                            with open(tf.name, "rb") as f:
                                resp = groq_client.audio.transcriptions.create(
                                    file=(chunk["name"], f.read()),
                                    model="whisper-large-v3",
                                    language=None if language == "auto" else language,
                                    response_format="text",
                                )
                            parts.append(f"[{short}]\n{resp}")
                            s.update(label=f"{short} — done ✓", state="complete")
                            break
                        except Exception as e:
                            if attempt < 2 and "rate" in str(e).lower():
                                wait_sec = 65
                                for remaining in range(wait_sec, 0, -1):
                                    s.update(
                                        label=f"Rate limit hit — retrying in {remaining}s …",
                                        state="running",
                                    )
                                    time.sleep(1)
                            else:
                                st.error(f"Transcription failed: {e}")
                                st.stop()
                finally:
                    os.unlink(tf.name)

            if i < len(selected) - 1:
                time.sleep(3)

        st.session_state.transcript = "\n\n".join(parts)

        # ── Generate documents via Claude (with prompt caching) ─
        claude     = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
        transcript = st.session_state.transcript

        # System prompt and transcript are cached across all 3 calls.
        # Only the task instruction differs per call.
        base_system = (
            f"You are a precise transcription analyst for a medical diagnostics company.\n"
            f"Company: {company}\n"
            f"Author: {sender_name}, {sender_title}\n"
            f"Context: {feedback_context}\n\n"
            f"ABSOLUTE RULES — apply to every output you produce:\n"
            f"1. Capture ONLY what was explicitly said in the transcript. Never infer, extrapolate, or add anything not directly stated.\n"
            f"2. Use the speaker's own words and phrasing as closely as possible. Do not paraphrase to make it sound better.\n"
            f"3. Do not add pleasantries, filler phrases, or padding sentences to meet a length target.\n"
            f"4. If a section has no content in the transcript, write exactly: [Not mentioned] — never leave it blank or fabricate content.\n"
            f"5. Do not validate, endorse, or editorialize the speaker's observations. Record them as stated.\n"
            f"6. Omit any section entirely if the transcript contains zero content for it."
        )

        def claude_call(task_prompt: str) -> str:
            msg = claude.messages.create(
                model="claude-sonnet-4-6",
                max_tokens=4000,
                system=[
                    {
                        "type": "text",
                        "text": base_system,
                        "cache_control": {"type": "ephemeral"},
                    }
                ],
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": f"TRANSCRIPT:\n\"\"\"\n{transcript}\n\"\"\"",
                                "cache_control": {"type": "ephemeral"},
                            },
                            {
                                "type": "text",
                                "text": task_prompt,
                            },
                        ],
                    }
                ],
                extra_headers={"anthropic-beta": "prompt-caching-2024-07-31"},
            )
            return msg.content[0].text

        if do_summary:
            with st.status("Generating summary …") as s:
                st.session_state.summary = claude_call(
                    """Write a factual bullet-point summary using only what was explicitly stated in the transcript above.

Rules:
- Use the speaker's own words and sentences as closely as possible.
- Each bullet = one clear statement directly from the transcript. No interpretation.
- Do NOT add context, background, or implications not stated by the speaker.
- Do NOT use phrases like "the speaker noted", "it was emphasized", "importantly" — just state the fact.
- If fewer than 5 distinct points were made, use fewer bullets. Do not pad.
- Maximum 8 bullets.

Format: plain bullet list, no heading, no preamble."""
                )
                s.update(label="Summary done ✓", state="complete")

        if do_mom:
            with st.status("Generating minutes of meeting …") as s:
                st.session_state.mom = claude_call(
                    f"""Generate Minutes of Meeting strictly from what was said in the transcript above.

RULES — non-negotiable:
1. Every line must trace directly to something said in the transcript. If it was not said, do not write it.
2. Quote or closely paraphrase the speaker's own words. Do not rewrite into polished prose.
3. For any section with no content in the transcript, write exactly: [Not mentioned in meeting] — never omit the section or fabricate content.
4. AGENDA: Only the host/initiator's stated purpose — verbatim or near-verbatim. No other participant's input here.
5. DISCUSSION: Attribute each point to the correct speaker using their actual words as closely as possible.
6. OBSERVATIONS & FINDINGS: Only observations explicitly stated. No inferences.
7. DECISIONS: Only decisions explicitly stated or agreed upon. No inferred decisions.
8. ACTION ITEMS: Only tasks explicitly assigned or volunteered. No derived tasks.

OUTPUT FORMAT — use exactly:

## Minutes of Meeting
**Date:** [exact date if stated; otherwise: Not mentioned]
**Prepared by:** {sender_name}, {sender_title}, {company}
**Attendees:** [names if mentioned; otherwise: Not mentioned]

### Agenda
[Host's stated purpose — verbatim or near-verbatim only]

### Discussion
[Each participant's points in their own words, attributed by name or role]

### Observations & Findings
[Only explicitly stated observations. If none: Not mentioned in meeting.]

### Decisions
[Only explicitly stated decisions. If none: Not mentioned in meeting.]

### Action Items
| # | Action | Owner | Due |
|---|--------|-------|-----|
[Only explicitly assigned tasks. If none: single row — Not discussed | — | —]"""
                )
                s.update(label="Minutes done ✓", state="complete")

        if do_email:
            with st.status("Generating feedback email …") as s:
                st.session_state.email = claude_call(
                    f"""Convert the transcript above into a professional feedback email.

Subject: {email_subject}
From: {sender_name}, {sender_title}, {company}
To: {recipient_name}

Rules:
- Include ONLY observations, issues, and points explicitly stated in the transcript.
- Use the speaker's own words and phrasing. Do not embellish or rewrite.
- Omit any section (observations / findings / next steps) if the transcript has no content for it.
- No padding, pleasantries beyond a single-line greeting, or filler sentences.
- Bullet points for issues and findings. One sentence per bullet.
- Close with sign-off from {sender_name}. No motivational or closing remarks beyond the sign-off."""
                )
                s.update(label="Email done ✓", state="complete")

        st.session_state.stage = "done"
        st.rerun()

# ═══════════════════════════════════════════════════════════════
# STAGE 3 — Outputs
# ═══════════════════════════════════════════════════════════════
elif st.session_state.stage == "done":
    st.success("Done!")
    stem       = st.session_state.stem
    summary    = st.session_state.summary
    mom        = st.session_state.mom
    email      = st.session_state.email
    transcript = st.session_state.transcript

    tab_labels = []
    if do_summary and summary: tab_labels.append("📋 Summary")
    if do_mom     and mom:     tab_labels.append("📝 Minutes of Meeting")
    if do_email   and email:   tab_labels.append("📧 Feedback Email")
    tab_labels.append("🗒️ Transcript")

    tabs = st.tabs(tab_labels)
    idx  = 0

    if do_summary and summary:
        with tabs[idx]:
            st.markdown(summary)
            c1, c2 = st.columns(2)
            c1.download_button("⬇ Download Summary (.md)", summary,
                               f"{stem}_summary.md", mime="text/markdown")
            c2.download_button("⬇ Download Summary (.docx)",
                               markdown_to_docx(summary, "Summary"),
                               f"{stem}_summary.docx", mime=DOCX_MIME)
        idx += 1

    if do_mom and mom:
        with tabs[idx]:
            st.markdown(mom)
            c1, c2 = st.columns(2)
            c1.download_button("⬇ Download MoM (.md)", mom,
                               f"{stem}_mom.md", mime="text/markdown")
            c2.download_button("⬇ Download MoM (.docx)",
                               markdown_to_docx(mom, "Minutes of Meeting"),
                               f"{stem}_mom.docx", mime=DOCX_MIME)
        idx += 1

    if do_email and email:
        with tabs[idx]:
            st.markdown(email)
            c1, c2 = st.columns(2)
            c1.download_button("⬇ Download Email (.md)", email,
                               f"{stem}_feedback_email.md", mime="text/markdown")
            c2.download_button("⬇ Download Email (.docx)",
                               markdown_to_docx(email, "Feedback Email"),
                               f"{stem}_feedback_email.docx", mime=DOCX_MIME)
        idx += 1

    with tabs[idx]:
        st.text_area("Raw transcript", transcript, height=300)
        st.download_button("⬇ Download Transcript", transcript,
                           f"{stem}_transcript.txt", mime="text/plain")

    st.divider()
    if st.button("↩ Process Another File"):
        for k, v in _defaults.items():
            st.session_state[k] = v
        st.rerun()
